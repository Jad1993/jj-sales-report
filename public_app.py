import streamlit as st
import csv
import io
from collections import defaultdict
import openpyxl

st.title("📊 JJ Daily Sales Report Tool")

# ---- Simple password gate ----
def check_password():
    def password_entered():
        if st.session_state["password"] == st.secrets["app_password"]:
            st.session_state["password_correct"] = True
            del st.session_state["password"]
        else:
            st.session_state["password_correct"] = False

    if "password_correct" not in st.session_state:
        st.text_input("Password", type="password", on_change=password_entered, key="password")
        return False
    elif not st.session_state["password_correct"]:
        st.text_input("Password", type="password", on_change=password_entered, key="password")
        st.error("Incorrect password")
        return False
    else:
        return True

if not check_password():
    st.stop()


BRAND_CELLS = {
    "ACER": {"daily": "B6", "utd": "B7", "unit_sold_start": "B13", "unit_sold_rows": 4},
    "HP":   {"daily": "G6", "utd": "G7", "unit_sold_start": "G13", "unit_sold_rows": 4},
    "MSI":  {"daily": "K6", "utd": "K7", "unit_sold_start": "K13", "unit_sold_rows": 4},
}

KNOWN_PREFIXES = ("NB-", "SW-", "CP-", "AIO-", "ADT-", "DC-", "CART-")


def clean_num(v):
    if not v:
        return "0"
    v = v.replace(",", "").replace("(", "-").replace(")", "")
    return v if v else "0"


def detect_brand(identifier):
    ident = identifier.upper()
    if "ACER" in ident:
        return "ACER"
    elif "MSI" in ident:
        return "MSI"
    elif "HP" in ident:
        return "HP"
    return None


def parse_multibranch_csv(file_obj):
    content = file_obj.read().decode("utf-8-sig")
    rows = list(csv.reader(io.StringIO(content)))

    header = rows[0]
    date_cols = {}
    col = 7
    while col < len(header):
        date_str = header[col].strip()
        if date_str:
            date_cols[date_str] = (col, col + 1, col + 2)
        col += 3

    TOTAL_AMOUNT_IDX = 5

    brand_data = {
        "ACER": {"total_amount": 0.0, "daily_by_date": defaultdict(float), "items": []},
        "HP": {"total_amount": 0.0, "daily_by_date": defaultdict(float), "items": []},
        "MSI": {"total_amount": 0.0, "daily_by_date": defaultdict(float), "items": []},
    }

    current_brand = None

    for row in rows[1:]:
        if not row or not row[0].strip():
            continue
        no_field = row[0].strip()

        if no_field == "GRAND TOTAL":
            break

        if "." not in no_field and no_field.replace("-", "").isdigit():
            branch_id = row[1].strip()
            brand = detect_brand(branch_id)
            current_brand = brand
            if brand:
                total_amt = float(clean_num(row[TOTAL_AMOUNT_IDX]))
                brand_data[brand]["total_amount"] += total_amt
                for date_str, (qcol, acol, gcol) in date_cols.items():
                    amt_str = row[acol] if acol < len(row) else ""
                    brand_data[brand]["daily_by_date"][date_str] += float(clean_num(amt_str))
        else:
            if current_brand:
                item_code = row[1]
                item_name = row[2]
                qty_by_date = {}
                for date_str, (qcol, acol, gcol) in date_cols.items():
                    q = row[qcol] if qcol < len(row) else ""
                    qty_by_date[date_str] = q.strip()
                brand_data[current_brand]["items"].append((item_code, item_name, qty_by_date))

    return brand_data, list(date_cols.keys())


def build_unit_sold_lines_for_date(items, selected_date):
    nb_brand_counts = defaultdict(int)
    sw_count = 0
    cp_count = 0
    aio_count = 0
    adt_count = 0
    dc_count = 0
    cart_count = 0

    for code, name, qty_by_date in items:
        q = qty_by_date.get(selected_date, "").strip()
        if not q:
            continue
        try:
            qty = int(float(q))
        except ValueError:
            continue
        if qty <= 0:
            continue

        if code.startswith("NB-"):
            parts = code.split("-")
            brand = parts[1] if len(parts) > 1 else "UNKNOWN"
            nb_brand_counts[brand] += qty
        elif code.startswith("SW-"):
            sw_count += qty
        elif code.startswith("CP-"):
            cp_count += qty
        elif code.startswith("AIO-"):
            aio_count += qty
        elif code.startswith("ADT-"):
            adt_count += qty
        elif code.startswith("DC-"):
            dc_count += qty
        elif code.startswith("CART-"):
            cart_count += qty

    lines = []
    for brand, qty in nb_brand_counts.items():
        lines.append(f"NB {brand} X{qty}")
    if sw_count > 0:
        lines.append(f"ANTI VIRUS X{sw_count}")
    if cp_count > 0:
        lines.append(f"AEW CP X{cp_count}")
    if aio_count > 0:
        lines.append(f"AIO X{aio_count}")
    if adt_count > 0:
        lines.append(f"ADT X{adt_count}")
    if dc_count > 0:
        lines.append(f"DC X{dc_count}")
    if cart_count > 0:
        lines.append(f"CART X{cart_count}")

    return lines


def build_other_items_summary(items, selected_date):
    """Groups every item NOT matching a known prefix, for review purposes."""
    prefix_counts = defaultdict(int)
    prefix_examples = {}

    for code, name, qty_by_date in items:
        if code.startswith(KNOWN_PREFIXES):
            continue

        q = qty_by_date.get(selected_date, "").strip()
        if not q:
            continue
        try:
            qty = int(float(q))
        except ValueError:
            continue
        if qty <= 0:
            continue

        prefix = code.split("-")[0] if "-" in code else code
        prefix_counts[prefix] += qty
        if prefix not in prefix_examples:
            prefix_examples[prefix] = name

    lines = []
    for prefix, qty in sorted(prefix_counts.items()):
        example = prefix_examples[prefix][:50]
        lines.append(f"{prefix}- X{qty}  (e.g. {example}...)")

    return lines


def write_branch_to_excel(wb, daily_cell, utd_cell, unit_sold_start_cell, unit_sold_rows,
                           daily_total, monthly_total, unit_sold_lines):
    ws = wb["Sheet1"]
    ws[daily_cell] = float(daily_total)
    ws[utd_cell] = float(monthly_total)

    start_col = unit_sold_start_cell[0]
    start_row = int(unit_sold_start_cell[1:])

    for r in range(unit_sold_rows):
        ws[f"{start_col}{start_row + r}"] = None

    for i, line in enumerate(unit_sold_lines):
        ws[f"{start_col}{start_row + i}"] = line


st.write("Upload your current Excel report and the combined multi-branch CSV export (JJ only), then review and save.")

excel_upload = st.file_uploader("Upload your Excel report (.xlsx)", type="xlsx", key="excel_upload")
csv_upload = st.file_uploader("Upload combined multi-branch CSV export (JJ only)", type="csv", key="csv_upload")

if excel_upload and csv_upload:
    wb = openpyxl.load_workbook(excel_upload)
    brand_data, all_dates = parse_multibranch_csv(csv_upload)

    def total_for_date(d):
        return sum(brand_data[b]["daily_by_date"].get(d, 0) for b in brand_data)

    dates_with_data = [d for d in all_dates if total_for_date(d) > 0]
    default_date = dates_with_data[-1] if dates_with_data else all_dates[-1]
    default_index = all_dates.index(default_date)

    selected_date = st.selectbox("Select report date", options=all_dates, index=default_index)

    if st.button("Preview All Branches"):
        results = {}
        for brand, cells in BRAND_CELLS.items():
            data = brand_data[brand]
            daily_total = data["daily_by_date"].get(selected_date, 0)
            utd_total = data["total_amount"]
            unit_sold_lines = build_unit_sold_lines_for_date(data["items"], selected_date)
            other_items = build_other_items_summary(data["items"], selected_date)

            results[brand] = {
                "daily_total": daily_total,
                "monthly_total": utd_total,
                "unit_sold_lines": unit_sold_lines
            }
            st.write(f"### {brand}")
            st.write(f"**DAILY ({selected_date}):** RM {daily_total:.2f}")
            st.write(f"**UTD:** RM {utd_total:.2f}")
            st.write("**Unit Sold (will be saved to Excel):**")
            if unit_sold_lines:
                for line in unit_sold_lines:
                    st.write(f"- {line}")
            else:
                st.write("- (none)")

            with st.expander(f"🔍 Other items sold ({brand}) — not yet in Unit Sold"):
                if other_items:
                    for line in other_items:
                        st.write(f"- {line}")
                    st.caption("Tell me which of these prefixes you want added to the official Unit Sold rule.")
                else:
                    st.write("(none — everything sold today is already covered)")

        st.session_state["pending_all"] = results

    if "pending_all" in st.session_state:
        st.warning("This will overwrite the ACER, HP, and MSI blocks in the Excel file.")
        if st.button("✅ Confirm & Save All"):
            results = st.session_state["pending_all"]
            for brand, cells in BRAND_CELLS.items():
                data = results[brand]
                write_branch_to_excel(
                    wb, cells["daily"], cells["utd"], cells["unit_sold_start"], cells["unit_sold_rows"],
                    data["daily_total"], data["monthly_total"], data["unit_sold_lines"]
                )
            output = io.BytesIO()
            wb.save(output)
            output.seek(0)
            st.success("All branches updated!")
            st.download_button(
                label="⬇️ Download Updated Excel File",
                data=output,
                file_name=f"JJ_SALES_REPORT_{selected_date}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
            del st.session_state["pending_all"]
else:
    st.info("Upload both the Excel report and the combined CSV to get started.")
